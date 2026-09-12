# -*- coding: utf-8 -*-
"""build_docx.py - 生成综测活动证明汇总 Word。

输入 plan JSON（由 agent 在归类、判定分值后生成）：
{
  "title": "本科生综合素质测评活动证明汇总",
  "student_name": "张三",
  "categories": [
    {
      "name": "一、思想品德",
      "entries": [
        {
          "activity": "XX 志愿服务活动",
          "score": 2,
          "note": "志愿时长 20 小时（可选备注，可空字符串）",
          "images": ["C:/abs/path/proof_p1.png", "C:/abs/path/proof_p2.png"]
        }
      ]
    }
  ]
}

输出 Word 结构：
  第 1 页：标题 + 学生信息 + 加分汇总表（含合计行）
  之后：按 categories 分章；每个 entry 的每张图片独占一页，
        页首为"活动名称（类别 | 分值：X 分）"标题，图片居中铺满版心。

用法：
  python build_docx.py --plan plan.json --out 综测证明汇总.docx
"""
import argparse
import json
import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PAGE_W, PAGE_H = Cm(21.0), Cm(29.7)   # A4
MARGIN = Cm(2.5)
MAX_IMG_W = PAGE_W - MARGIN * 2        # 16 cm
MAX_IMG_H = PAGE_H - MARGIN * 2 - Cm(1.8)  # 预留标题行高度


def set_font(run, cn="宋体", en="Times New Roman", size=None, bold=None, color=None):
    run.font.name = en
    r = run._element.rPr
    r.rFonts.set(qn("w:eastAsia"), cn)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_par(doc, text, size=12, bold=False, align=None, cn="宋体", color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_font(run, cn=cn, size=size, bold=bold, color=color)
    return p


def img_size(path):
    from PIL import Image
    with Image.open(path) as im:
        return im.size  # (w, h) px


def add_image_fit(doc, path):
    w_px, h_px = img_size(path)
    max_w_emu, max_h_emu = int(MAX_IMG_W), int(MAX_IMG_H)
    # 按像素比例换算：先按宽度缩放，超高度则按高度缩放
    scale = min(max_w_emu / w_px, max_h_emu / h_px)
    w, h = int(w_px * scale), int(h_px * scale)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(path, width=w, height=h)


def fmt_score(v):
    if isinstance(v, (int, float)):
        return f"{v:g}"
    return str(v)


def numeric(v):
    try:
        return float(v)
    except (TypeError, ValueError):
        return 0.0


def build(plan_path, out_path):
    with open(plan_path, "r", encoding="utf-8") as fp:
        plan = json.load(fp)

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = PAGE_W, PAGE_H
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = MARGIN

    # ---- 封面区：标题 + 信息 ----
    add_par(doc, plan.get("title", "综合素质测评活动证明汇总"),
            size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, cn="黑体")
    info = f"姓名：{plan.get('student_name', '')}"
    if plan.get("student_id"):
        info += f"    学号：{plan['student_id']}"
    add_par(doc, info, size=12, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_par(doc, "", size=6)

    # ---- 汇总表 ----
    add_par(doc, "加分汇总表", size=14, bold=True, cn="黑体")
    entries_flat = []
    for cat in plan.get("categories", []):
        for e in cat.get("entries", []):
            entries_flat.append((cat["name"], e))

    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["序号", "活动名称", "所属类别", "分值", "备注"]
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_font(run, cn="黑体", size=10.5, bold=True)

    total = 0.0
    for idx, (cat_name, e) in enumerate(entries_flat, 1):
        total += numeric(e.get("score"))
        row = table.add_row().cells
        vals = [str(idx), e.get("activity", ""), cat_name,
                fmt_score(e.get("score", "")), e.get("note", "")]
        for c, v in zip(row, vals):
            c.text = ""
            run = c.paragraphs[0].add_run(v)
            set_font(run, size=10.5)

    row = table.add_row().cells
    row[0].merge(row[1]).merge(row[2])
    run = row[0].paragraphs[0].add_run("合计")
    set_font(run, cn="黑体", size=10.5, bold=True)
    run = row[3].paragraphs[0].add_run(f"{total:g}")
    set_font(run, cn="黑体", size=10.5, bold=True)
    row[4].text = ""

    # ---- 正文：按类别分章，一页一图 ----
    first_block = True
    for cat in plan.get("categories", []):
        if not cat.get("entries"):
            continue
        doc.add_page_break()
        add_par(doc, cat["name"], size=15, bold=True, cn="黑体")
        first_block = False
        for e in cat["entries"]:
            for i, img in enumerate(e.get("images", [])):
                doc.add_page_break()
                head = f"{e.get('activity', '')}（{cat['name']} | 分值：{fmt_score(e.get('score', ''))}）"
                if len(e.get("images", [])) > 1:
                    head += f"  第 {i + 1}/{len(e['images'])} 页"
                add_par(doc, head, size=12, bold=True, cn="黑体")
                if not os.path.isfile(img):
                    add_par(doc, f"[缺失图片：{img}]", size=10.5, color=(0xC0, 0, 0))
                    continue
                add_image_fit(doc, img)

    doc.save(out_path)
    print(f"saved: {out_path}")
    print(f"categories={len(plan.get('categories', []))}, entries={len(entries_flat)}, total_score={total:g}")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--plan", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()
    build(args.plan, args.out)


if __name__ == "__main__":
    main()
